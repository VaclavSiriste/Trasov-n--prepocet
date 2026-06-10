#!/usr/bin/env python3
"""Import Raynet listů z Excelu do PostgreSQL."""

import argparse
import re
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / ".pylibs"))
    import openpyxl

try:
    import psycopg2
    from psycopg2.extras import execute_batch
except ImportError:
    psycopg2 = None

SHEET_CONFIG = {
    "raynet_montaze": {
        "sheet": "Raynet MONTÁŽE",
        "table": "raynet_montaze",
        "columns": [
            "kategorie", "naplanovano_od", "naplanovano_do", "trvani", "predmet",
            "ucastnici", "monteri", "monter_c_1", "monter_c_2", "monter_c_3",
            "monteru", "pocet_monterohodin", "hodin", "mesic", "naplanovano_od_datum",
            "mesic_datum", "misto_setkani", "kraj", "rok", "stitky",
        ],
        "col_count": 20,
    },
    "raynet_zamerovaci": {
        "sheet": "Raynet ZAMĚŘOVAČI",
        "table": "raynet_zamerovaci",
        "columns": [
            "kategorie", "naplanovano_od", "naplanovano_do", "trvani", "predmet",
            "ucastnici", "zamerovac", "zamerovac_c_1", "zamerovac_c_2", "zamerovac_c_3",
            "zamerovacu", "pocet_hodin_zamereni", "hodin", "mesic", "naplanovano_od_datum",
            "mesic_datum", "misto_setkani", "kraj", "rok", "stitky",
        ],
        "col_count": 20,
    },
    "raynet_obvolat_rma": {
        "sheet": "Raynet OBVOLAT RMA",
        "table": "raynet_obvolat_rma",
        "columns": [
            "kategorie", "naplanovano_od", "naplanovano_do", "trvani", "predmet", "ucastnici",
        ],
        "col_count": 6,
    },
    "raynet_prejezdy": {
        "sheet": "Raynet PREJEZDY",
        "table": "raynet_prejezdy",
        "columns": [
            "kategorie", "cas_od", "cas_do", "trvani", "predmet", "technik", "hodiny",
            "naplanovano_od", "naplanovano_do", "mesic", "rok", "tyden", "mvt",
            "pocet_km_na_zakazku",
        ],
        "col_count": 14,
    },
}


def to_int(val):
    if val is None or val == "":
        return None
    return int(float(val))


def to_numeric(val):
    if val is None or val == "":
        return None
    return float(val)


def to_text(val):
    if val is None:
        return None
    s = str(val).strip()
    return s or None


def to_timestamptz(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val
    return None


def to_date(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def to_interval(val):
    if val is None or val == "":
        return None
    if isinstance(val, timedelta):
        return val
    return None


def to_time(val):
    if val is None or val == "":
        return None
    if isinstance(val, time):
        return val
    if isinstance(val, timedelta):
        total = int(val.total_seconds())
        h, rem = divmod(total, 3600)
        m, s = divmod(rem, 60)
        return time(h, m, s)
    return None


CONVERTERS = {
    "kategorie": to_text,
    "predmet": to_text,
    "ucastnici": to_text,
    "monteri": to_text,
    "monter_c_1": to_text,
    "monter_c_2": to_text,
    "monter_c_3": to_text,
    "zamerovac": to_text,
    "zamerovac_c_1": to_text,
    "zamerovac_c_2": to_text,
    "zamerovac_c_3": to_text,
    "misto_setkani": to_text,
    "kraj": to_text,
    "stitky": to_text,
    "technik": to_text,
    "mvt": to_text,
    "naplanovano_od": to_timestamptz,
    "naplanovano_do": to_timestamptz,
    "cas_od": to_timestamptz,
    "cas_do": to_timestamptz,
    "naplanovano_od_datum": to_date,
    "trvani": to_interval,
    "monteru": to_numeric,
    "pocet_monterohodin": to_numeric,
    "hodin": to_numeric,
    "zamerovacu": to_numeric,
    "pocet_hodin_zamereni": to_numeric,
    "hodiny": to_numeric,
    "predmet_prejezdy": to_numeric,
    "pocet_km_na_zakazku": to_numeric,
    "mesic": to_int,
    "mesic_datum": to_int,
    "rok": to_int,
    "tyden": to_int,
}


FIELD_TYPES = {
    "raynet_montaze": {
        "naplanovano_od": "ts", "naplanovano_do": "ts", "trvani": "interval",
        "monteru": "num", "pocet_monterohodin": "num", "hodin": "num",
        "mesic": "int", "naplanovano_od_datum": "date", "mesic_datum": "int", "rok": "int",
    },
    "raynet_zamerovaci": {
        "naplanovano_od": "ts", "naplanovano_do": "ts", "trvani": "interval",
        "zamerovacu": "num", "pocet_hodin_zamereni": "num", "hodin": "num",
        "mesic": "int", "naplanovano_od_datum": "date", "mesic_datum": "int", "rok": "int",
    },
    "raynet_obvolat_rma": {
        "naplanovano_od": "ts", "naplanovano_do": "ts", "trvani": "interval",
    },
    "raynet_prejezdy": {
        "cas_od": "ts", "cas_do": "ts", "trvani": "time", "predmet": "num",
        "hodiny": "num", "naplanovano_od": "date", "naplanovano_do": "date",
        "mesic": "int", "rok": "int", "tyden": "int", "pocet_km_na_zakazku": "num",
    },
}


def convert_value(kind, val):
    if kind == "ts":
        return to_timestamptz(val)
    if kind == "date":
        return to_date(val)
    if kind == "interval":
        return to_interval(val)
    if kind == "time":
        return to_time(val)
    if kind == "int":
        return to_int(val)
    if kind == "num":
        return to_numeric(val)
    return to_text(val)


def convert_row(table_key, columns, raw):
    types = FIELD_TYPES.get(table_key, {})
    return tuple(convert_value(types.get(col, "text"), val) for col, val in zip(columns, raw))


def read_sheet_rows(wb, table_key, config):
    ws = wb[config["sheet"]]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        raw = list(row[: config["col_count"]])
        while len(raw) < config["col_count"]:
            raw.append(None)
        rows.append(convert_row(table_key, config["columns"], raw))
    return rows


def import_all(xlsx_path, dsn, truncate=False):
    if psycopg2 is None:
        raise RuntimeError("Nainstalujte psycopg2: pip install psycopg2-binary")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    conn = psycopg2.connect(dsn)
    conn.autocommit = False

    try:
        with conn.cursor() as cur:
            for key, config in SHEET_CONFIG.items():
                rows = read_sheet_rows(wb, key, config)
                if truncate:
                    cur.execute(f"TRUNCATE TABLE {config['table']} RESTART IDENTITY")
                cols = ", ".join(config["columns"])
                placeholders = ", ".join(["%s"] * len(config["columns"]))
                sql = f"INSERT INTO {config['table']} ({cols}) VALUES ({placeholders})"
                execute_batch(cur, sql, rows, page_size=500)
                print(f"{config['table']}: {len(rows)} řádků")
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def main():
    parser = argparse.ArgumentParser(description="Import Raynet Excel -> PostgreSQL")
    parser.add_argument(
        "xlsx",
        nargs="?",
        default="/Users/vaclavsiriste/Downloads/Kopie souboru Trasování (1).xlsx",
    )
    parser.add_argument(
        "--dsn",
        default="postgresql://localhost/trasovani",
        help="PostgreSQL connection string",
    )
    parser.add_argument("--truncate", action="store_true", help="Smazat data před importem")
    args = parser.parse_args()
    import_all(args.xlsx, args.dsn, truncate=args.truncate)


if __name__ == "__main__":
    main()
